"""Estrazione testo da PDF / Word / testo. Deterministico, nessun LLM."""
from pathlib import Path

TEXT_EXT = {".txt", ".md", ".csv", ".log", ".json"}
SUPPORTED = {".pdf", ".docx"} | TEXT_EXT


class DocumentError(ValueError):
    pass


def extract_text(path: Path, max_chars: int) -> tuple[str, bool]:
    """Ritorna (testo, troncato)."""
    ext = path.suffix.lower()
    if ext not in SUPPORTED:
        raise DocumentError(f"Formato non supportato ({ext or 'nessuna estensione'}). Ammessi: PDF, DOCX, TXT/MD/CSV/JSON/LOG.")
    try:
        if ext == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(str(path))
            if reader.is_encrypted:
                raise DocumentError("PDF protetto da password.")
            parts, size = [], 0
            for page in reader.pages:
                t = page.extract_text() or ""
                parts.append(t)
                size += len(t)
                if size > max_chars:
                    break
            text = "\n\n".join(parts)
        elif ext == ".docx":
            from docx import Document

            doc = Document(str(path))
            text = "\n".join(p.text for p in doc.paragraphs)
            for table in doc.tables:
                for row in table.rows:
                    text += "\n" + " | ".join(c.text.strip() for c in row.cells)
        else:
            text = path.read_text(encoding="utf-8", errors="replace")
    except DocumentError:
        raise
    except Exception as e:  # file corrotto
        raise DocumentError(f"Impossibile leggere il file: {type(e).__name__}") from e

    text = text.strip()
    if not text:
        raise DocumentError("Nessun testo estraibile (PDF scansionato? serve OCR).")
    return (text[:max_chars], True) if len(text) > max_chars else (text, False)


def split_chunks(text: str, size: int) -> list[str]:
    chunks, cur = [], ""
    for para in text.split("\n"):
        if len(cur) + len(para) + 1 > size and cur:
            chunks.append(cur)
            cur = ""
        while len(para) > size:  # riga enorme
            chunks.append(para[:size])
            para = para[size:]
        cur += para + "\n"
    if cur.strip():
        chunks.append(cur)
    return chunks
